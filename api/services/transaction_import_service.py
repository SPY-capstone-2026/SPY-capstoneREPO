"""
거래내역 파일 가져오기 (.xlsx / .csv 모두).

파일 형식은 확장자가 아니라 '열 이름'으로 판별한다.

1) TOSS    : 토스뱅크 앱에서 받은 거래내역
   열: 거래 일시 | 적요 | 거래 유형 | 거래 기관 | 계좌번호 | 거래 금액 | 거래 후 잔액 | 메모
   위쪽에 성명/조회기간/안내 문구가 있고 '거래 일시' 줄부터 표가 시작된다. 출금이 음수.

2) GENERIC : 사용자가 직접 만든 표 (기존 프론트 CSV 업로드 형식과 같은 열 이름)
   필수: 날짜 | 금액 | 결제처      선택: 카테고리 | 시간 | 메모
   모든 행을 지출로 본다. 금액 부호는 파일 전체에서 다수인 쪽을 지출로 보고, 반대 부호는 수입으로 제외한다.

행 판정 (kind) — 저장 가능한 건 EXPENSE, TRANSFER, OTHER_OUT 뿐이다.
- EXPENSE          : 지출 (토스 카드 결제 / 직접 만든 표의 일반 행)  → 기본 포함
- REFUND           : 토스 카드 결제 취소(양수)                     → 저장 안 함, 짝이 되는 결제 상쇄
- TRANSFER         : 토스 다른 계좌로 보낸 출금                    → 기본 제외 (사용자가 켤 수 있음)
- BENEFIT_REVERSAL : 토스 캐시백·복권·이자 등 혜택 회수           → 제외
- OTHER_OUT        : 토스 그 밖의 출금(자동이체 등)               → 추천 카테고리가 있으면 기본 포함
- INCOME           : 입금/수입                                     → 제외 (수입은 저장하지 않는다)
- INVALID          : 날짜·금액·결제처를 읽을 수 없는 행           → 제외, error 에 이유

수입을 저장하지 않으므로 정산·리포트·AI 의 '지출 = amount 합계' 규칙은 그대로 유지된다.
"""
import csv
import hashlib
import io
import re
from collections import defaultdict
from datetime import date, datetime, time, timedelta
from typing import Dict, List, Optional, Tuple

from fastapi import HTTPException
from sqlalchemy import delete
from sqlalchemy.exc import IntegrityError
from sqlmodel import Session, select

from models import GeneratedReport, Transaction
from services.category_classifier import (
    classify_merchants,
    get_user_categories,
    merchant_key,
    normalize_merchant,
    remember_user_choice,
)

FORMAT_TOSS = "TOSS"
FORMAT_GENERIC = "GENERIC"

KIND_EXPENSE = "EXPENSE"
KIND_REFUND = "REFUND"
KIND_TRANSFER = "TRANSFER"
KIND_BENEFIT_REVERSAL = "BENEFIT_REVERSAL"
KIND_OTHER_OUT = "OTHER_OUT"
KIND_INCOME = "INCOME"
KIND_INVALID = "INVALID"
IMPORTABLE_KINDS = (KIND_EXPENSE, KIND_TRANSFER, KIND_OTHER_OUT)

SOURCE_FILE = "FILE"  # 직접 만든 표의 '카테고리' 열 값을 그대로 쓴 경우
LABEL_TOSS = "토스뱅크"
LABEL_GENERIC = "파일 업로드"

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_ROWS = 10_000
HEADER_SEARCH_ROWS = 30
REFUND_LOOKBACK_DAYS = 90
DEFAULT_TIME = time(12, 0)  # 직접 만든 표에 시간이 없을 때

TOSS_REQUIRED = ("거래 일시", "적요", "거래 유형", "거래 금액")
GENERIC_ALIASES = {
    "date": ("날짜", "거래일", "거래일자", "date", "tx_date"),
    "amount": ("금액", "거래금액", "amount"),
    "merchant": ("결제처", "사용처", "가맹점", "가맹점명", "merchant", "merchant_name"),
    "category": ("카테고리", "분류", "category", "final_category"),
    "time": ("시간", "거래시간", "time", "tx_time"),
    "memo": ("메모", "memo"),
}
_BENEFIT_RE = re.compile(r"캐시백|복권|이자|포인트|리워드|적립")


# ----------------------------------------
# 파일 → 표
# ----------------------------------------
def read_table(content: bytes, filename: str) -> List[list]:
    name = (filename or "").lower()
    if len(content) > MAX_FILE_BYTES:
        raise HTTPException(status_code=400, detail="파일이 너무 큽니다 (최대 5MB). 기간을 나눠서 올려 주세요")

    if name.endswith(".csv"):
        for encoding in ("utf-8-sig", "cp949"):
            try:
                text = content.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise HTTPException(status_code=400, detail="CSV 인코딩을 읽을 수 없습니다 (UTF-8 또는 CP949로 저장해 주세요)")
        try:
            return [row for row in csv.reader(io.StringIO(text))]
        except csv.Error:
            raise HTTPException(status_code=400, detail="CSV 형식이 올바르지 않습니다 (따옴표가 닫히지 않았는지 확인해 주세요)")

    if name.endswith(".xlsx"):
        from openpyxl import load_workbook
        try:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception:
            raise HTTPException(
                status_code=400,
                detail="엑셀 파일을 열 수 없습니다. 암호가 걸려 있다면 암호를 해제해 다시 저장해 주세요",
            )
        rows = [list(r) for r in workbook.worksheets[0].iter_rows(values_only=True)]
        workbook.close()
        return rows

    raise HTTPException(status_code=400, detail=".xlsx 또는 .csv 파일만 올릴 수 있습니다")


def _cell(value) -> str:
    return "" if value is None else str(value).strip()


def _norm_header(value) -> str:
    return _cell(value).replace(" ", "").lower()


def detect_format(table: List[list]) -> Tuple[str, int, Dict[str, int]]:
    """반환: (형식, 머리글 줄 index, {열 이름: 열 index})"""
    toss_keys = {_norm_header(h): h for h in TOSS_REQUIRED}
    for i, row in enumerate(table[:HEADER_SEARCH_ROWS]):
        cells = [_norm_header(c) for c in row]
        if all(k in cells for k in toss_keys):
            return FORMAT_TOSS, i, {_cell(c): j for j, c in enumerate(row) if _cell(c)}
        found = {}
        for field, aliases in GENERIC_ALIASES.items():
            for j, c in enumerate(cells):
                if c in {_norm_header(a) for a in aliases}:
                    found[field] = j
                    break
        if all(f in found for f in ("date", "amount", "merchant")):
            return FORMAT_GENERIC, i, found
    raise HTTPException(
        status_code=400,
        detail="거래내역 형식을 알 수 없습니다. 토스뱅크에서 받은 파일이거나, "
               "첫 줄에 '날짜, 금액, 결제처' 열(선택: 카테고리, 시간, 메모)이 있어야 합니다",
    )


# ----------------------------------------
# 값 읽기
# ----------------------------------------
def _parse_amount(value) -> int:
    if isinstance(value, (int, float)):
        return int(round(value))
    text = re.sub(r"[₩원,\s]", "", _cell(value))
    if not text:
        raise ValueError("empty")
    return int(round(float(text)))


_DATE_FORMATS = ("%Y-%m-%d", "%Y.%m.%d", "%Y/%m/%d", "%Y%m%d")
_TIME_FORMATS = ("%H:%M:%S", "%H:%M")


def _parse_date_time(value) -> Tuple[date, Optional[time]]:
    """날짜 칸에 '2026-10-29' 또는 '2026.10.29 14:57:48' 같은 날짜+시간이 올 수 있다."""
    if isinstance(value, datetime):
        t = value.time().replace(microsecond=0)
        return value.date(), (t if t != time(0, 0) else None)
    if isinstance(value, date):
        return value, None
    text = _cell(value).rstrip(".")
    parts = text.split()
    for fmt in _DATE_FORMATS:
        try:
            d = datetime.strptime(parts[0], fmt).date()
            break
        except (ValueError, IndexError):
            continue
    else:
        raise ValueError(text)
    t = _parse_time(parts[1]) if len(parts) > 1 else None
    return d, t


def _parse_time(value) -> Optional[time]:
    if value is None or _cell(value) == "":
        return None
    if isinstance(value, time):
        return value.replace(microsecond=0)
    if isinstance(value, datetime):
        return value.time().replace(microsecond=0)
    text = _cell(value)
    for fmt in _TIME_FORMATS:
        try:
            return datetime.strptime(text, fmt).time()
        except ValueError:
            continue
    raise ValueError(text)


def _check_size(rows: list):
    if len(rows) > MAX_ROWS:
        raise HTTPException(status_code=400, detail=f"한 번에 {MAX_ROWS:,}건까지 올릴 수 있습니다. 기간을 나눠 주세요")


# ----------------------------------------
# 형식별 행 읽기 → 공통 행 dict
# 공통 키: row_no, occurred_at(datetime|None), summary, merchant_name, amount(양수), kind, error,
#          tx_type, memo, file_category, key_parts
# ----------------------------------------
def _toss_kind(amount: int, tx_type: str, summary: str, institution: str, has_account: bool) -> str:
    is_card = "카드" in tx_type or "결제" in tx_type
    if amount < 0:
        if is_card:
            return KIND_EXPENSE
        if _BENEFIT_RE.search(summary):
            return KIND_BENEFIT_REVERSAL
        if institution or has_account:
            return KIND_TRANSFER
        return KIND_OTHER_OUT
    if amount > 0 and is_card:
        return KIND_REFUND
    return KIND_INCOME


def parse_toss_rows(table: List[list], header_index: int, col: Dict[str, int]) -> List[dict]:
    def get(row, name):
        idx = col.get(name)
        return row[idx] if idx is not None and idx < len(row) else None

    out = []
    for row_no, row in enumerate(table[header_index + 1:], start=header_index + 2):
        if not any(_cell(c) for c in row):
            continue
        summary = _cell(get(row, "적요"))
        tx_type = _cell(get(row, "거래 유형"))
        base = {"row_no": row_no, "summary": summary, "merchant_name": normalize_merchant(summary),
                "tx_type": tx_type, "memo": _cell(get(row, "메모"))[:200] or None,
                "file_category": None, "occurred_at": None, "amount": 0, "error": None}
        try:
            d, t = _parse_date_time(get(row, "거래 일시"))
            signed = _parse_amount(get(row, "거래 금액"))
        except (ValueError, TypeError):
            out.append({**base, "kind": KIND_INVALID, "error": "거래 일시나 거래 금액을 읽을 수 없습니다", "key_parts": None})
            continue
        try:
            balance = _parse_amount(get(row, "거래 후 잔액"))
        except (ValueError, TypeError):
            balance = None  # 잔액은 중복 판별 보조용이라 못 읽어도 진행
        if signed == 0:
            out.append({**base, "kind": KIND_INVALID, "error": "금액이 0원입니다", "key_parts": None})
            continue
        occurred_at = datetime.combine(d, t or time(0, 0))
        kind = _toss_kind(signed, tx_type, summary, _cell(get(row, "거래 기관")), bool(_cell(get(row, "계좌번호"))))
        out.append({**base, "occurred_at": occurred_at, "amount": abs(signed), "signed": signed, "kind": kind,
                    "key_parts": [FORMAT_TOSS, occurred_at.isoformat(), summary, str(signed),
                                  "" if balance is None else str(balance)]})
        _check_size(out)
    return out


def parse_generic_rows(table: List[list], header_index: int, col: Dict[str, int]) -> List[dict]:
    def get(row, field):
        idx = col.get(field)
        return row[idx] if idx is not None and idx < len(row) else None

    parsed = []
    for row_no, row in enumerate(table[header_index + 1:], start=header_index + 2):
        if not any(_cell(c) for c in row):
            continue
        merchant = _cell(get(row, "merchant"))
        base = {"row_no": row_no, "summary": merchant, "merchant_name": normalize_merchant(merchant),
                "tx_type": None, "memo": _cell(get(row, "memo"))[:200] or None,
                "file_category": _cell(get(row, "category")) or None,
                "occurred_at": None, "amount": 0, "error": None, "key_parts": None}
        try:
            d, t_from_date = _parse_date_time(get(row, "date"))
        except (ValueError, TypeError):
            parsed.append({**base, "kind": KIND_INVALID, "error": "날짜는 YYYY-MM-DD 형식으로 입력해 주세요"})
            continue
        try:
            t = _parse_time(get(row, "time")) or t_from_date
        except (ValueError, TypeError):
            parsed.append({**base, "kind": KIND_INVALID, "error": "시간은 HH:MM 형식으로 입력해 주세요"})
            continue
        try:
            signed = _parse_amount(get(row, "amount"))
        except (ValueError, TypeError):
            parsed.append({**base, "kind": KIND_INVALID, "error": "금액은 숫자로 입력해 주세요"})
            continue
        if signed == 0:
            parsed.append({**base, "kind": KIND_INVALID, "error": "금액이 0원입니다"})
            continue
        if not merchant:
            parsed.append({**base, "kind": KIND_INVALID, "error": "결제처가 비어 있습니다"})
            continue
        parsed.append({**base, "occurred_at": datetime.combine(d, t or DEFAULT_TIME), "amount": abs(signed),
                       "signed": signed, "kind": None, "time_given": t is not None})
        _check_size(parsed)

    # 부호: 파일에서 더 많은 쪽을 지출로 본다 (대부분 양수로 적지만, 은행 표를 옮기면 음수일 수 있음)
    valid = [r for r in parsed if r["kind"] is None]
    negatives = sum(1 for r in valid if r["signed"] < 0)
    spend_sign = -1 if negatives > len(valid) - negatives else 1

    seen = defaultdict(int)
    for r in valid:
        r["kind"] = KIND_EXPENSE if (r["signed"] > 0) == (spend_sign > 0) else KIND_INCOME
        # 같은 날·같은 시각·같은 결제처·같은 금액이 여러 줄이면 각각 다른 거래로 본다 (n번째)
        ident = (r["occurred_at"].isoformat(), r["summary"], r["signed"])
        seen[ident] += 1
        r["key_parts"] = [FORMAT_GENERIC, *map(str, ident), str(seen[ident])]
    return parsed


def parse_file(content: bytes, filename: str) -> Tuple[str, List[dict]]:
    table = read_table(content, filename)
    fmt, header_index, col = detect_format(table)
    rows = parse_toss_rows(table, header_index, col) if fmt == FORMAT_TOSS else parse_generic_rows(table, header_index, col)
    _check_size(rows)
    return fmt, rows


def make_import_key(user_id: str, key_parts: List[str]) -> str:
    return hashlib.sha256("|".join([user_id, *key_parts]).encode()).hexdigest()[:32]


# ----------------------------------------
# 미리보기
# ----------------------------------------
def _existing_keys(session: Session, user_id: str, keys: List[str]) -> set:
    found = set()
    for i in range(0, len(keys), 500):
        found |= set(session.exec(
            select(Transaction.import_key)
            .where(Transaction.user_id == user_id)
            .where(Transaction.import_key.in_(keys[i:i + 500]))
        ).all())
    return found


def _find_imported_expense(session: Session, user_id: str, refund: dict) -> Optional[Transaction]:
    refund_date = refund["occurred_at"].date()
    candidates = session.exec(
        select(Transaction)
        .where(Transaction.user_id == user_id)
        .where(Transaction.import_key.is_not(None))
        .where(Transaction.amount == refund["amount"])
        .where(Transaction.tx_date <= refund_date)
        .where(Transaction.tx_date >= refund_date - timedelta(days=REFUND_LOOKBACK_DAYS))
        .order_by(Transaction.tx_date.desc(), Transaction.tx_time.desc())
    ).all()
    key = merchant_key(refund["summary"])
    return next((tx for tx in candidates if merchant_key(tx.merchant_name) == key), None)


def _pair_refunds(session: Session, user_id: str, rows: List[dict], existing: set) -> None:
    """토스 환불 ↔ 결제: 같은 거래처·같은 금액, 환불보다 이전 결제 중 가장 가까운 것."""
    expenses = defaultdict(list)
    for r in rows:
        if r["kind"] == KIND_EXPENSE:
            expenses[(merchant_key(r["summary"]), r["amount"])].append(r)
    used = set()
    for refund in sorted((r for r in rows if r["kind"] == KIND_REFUND), key=lambda r: r["occurred_at"]):
        candidates = [e for e in expenses[(merchant_key(refund["summary"]), refund["amount"])]
                      if e["occurred_at"] <= refund["occurred_at"] and e["row_no"] not in used]
        if candidates:
            target = max(candidates, key=lambda e: e["occurred_at"])
            used.add(target["row_no"])
            target["refunded_by_row"] = refund["row_no"]
            refund["refunds_row"] = target["row_no"]
            if target["import_key"] in existing:
                # 결제는 예전에 이미 가져왔다 → DB 에 있는 그 거래를 취소 대상으로
                tx = session.exec(
                    select(Transaction)
                    .where(Transaction.user_id == user_id)
                    .where(Transaction.import_key == target["import_key"])
                ).first()
                if tx is not None:
                    refund["refunds_tx_id"] = tx.tx_id
            continue
        # 파일 안에 없으면 예전에 가져온 거래에서 찾는다 (직접 입력/시드 거래는 건드리지 않음)
        match = _find_imported_expense(session, user_id, refund)
        if match is not None:
            refund["refunds_tx_id"] = match.tx_id


def build_preview(session: Session, user_id: str, fmt: str, rows: List[dict]) -> dict:
    for r in rows:
        r["import_key"] = make_import_key(user_id, r["key_parts"]) if r["key_parts"] else None

    existing = _existing_keys(session, user_id, [r["import_key"] for r in rows if r["import_key"]])
    if fmt == FORMAT_TOSS:
        _pair_refunds(session, user_id, rows, existing)

    user_categories = get_user_categories(session, user_id)
    owned = set(user_categories)
    suggestions = classify_merchants(session, user_id, {r["summary"] for r in rows if r["summary"]})

    items, counts = [], defaultdict(int)
    for r in rows:
        category, source = suggestions.get(r["summary"], (None, "NONE"))
        warning = None
        if r["file_category"]:
            if r["file_category"] in owned:
                category, source = r["file_category"], SOURCE_FILE
            else:
                warning = f"'{r['file_category']}'은(는) 등록되지 않은 카테고리라 추천 카테고리로 바꿨어요"

        duplicate = bool(r["import_key"]) and r["import_key"] in existing
        refunded = "refunded_by_row" in r
        importable = r["kind"] in IMPORTABLE_KINDS and not duplicate and not refunded
        if r["kind"] == KIND_EXPENSE:
            include = importable
        elif r["kind"] == KIND_OTHER_OUT:
            include = importable and category is not None
        else:
            include = False

        counts[r["kind"]] += 1
        counts["duplicate"] += int(duplicate)
        counts["included"] += int(include)
        items.append({
            "row_no": r["row_no"],
            "import_key": r["import_key"],
            "occurred_at": r["occurred_at"].isoformat() if r["occurred_at"] else None,
            "kind": r["kind"],
            "tx_type": r["tx_type"],
            "summary": r["summary"],
            "merchant_name": r["merchant_name"],
            "amount": r["amount"],
            "memo": r["memo"],
            "suggested_category": category,
            "suggestion_source": source,
            "importable": importable,
            "include": include,
            "duplicate": duplicate,
            "error": r["error"],
            "warning": warning,
            "refunded_by_row": r.get("refunded_by_row"),
            "refunds_row": r.get("refunds_row"),
            "refunds_tx_id": r.get("refunds_tx_id"),
        })

    return {
        "source_format": fmt,
        "categories": user_categories,
        "summary": {
            "total_rows": len(rows),
            "included": counts["included"],
            "expense": counts[KIND_EXPENSE],
            "refund": counts[KIND_REFUND],
            "transfer": counts[KIND_TRANSFER],
            "other_out": counts[KIND_OTHER_OUT],
            "income": counts[KIND_INCOME],
            "benefit_reversal": counts[KIND_BENEFIT_REVERSAL],
            "invalid": counts[KIND_INVALID],
            "duplicate": counts["duplicate"],
        },
        "rows": items,
    }


# ----------------------------------------
# 등록 (미리보기에서 사용자가 확정한 행)
# ----------------------------------------
def confirm_import(session: Session, user_id: str, source_format: str, rows: List[dict],
                   cancel_tx_ids: List[str]) -> dict:
    """
    rows: [{import_key, occurred_at, merchant_name, amount(양수), final_category, memo?, tx_type?}]
    cancel_tx_ids: 미리보기의 refunds_tx_id (예전에 가져온 결제 중 이번 파일에서 취소된 것)
    전부 아니면 전무: 한 행이라도 잘못되면 아무것도 저장하지 않는다. 이미 있는 import_key 는 건너뛴다.
    """
    if source_format not in (FORMAT_TOSS, FORMAT_GENERIC):
        raise HTTPException(status_code=400, detail="source_format 은 TOSS 또는 GENERIC 입니다")
    if len(rows) > MAX_ROWS:
        raise HTTPException(status_code=400, detail=f"한 번에 {MAX_ROWS:,}건까지 등록할 수 있습니다")

    user_categories = set(get_user_categories(session, user_id))
    seen, prepared = set(), []
    for i, row in enumerate(rows, start=1):
        if row["import_key"] in seen:
            continue
        seen.add(row["import_key"])
        if row["amount"] <= 0:
            raise HTTPException(status_code=400, detail=f"{i}번째 항목: 금액은 0보다 커야 합니다")
        if not row.get("merchant_name", "").strip():
            raise HTTPException(status_code=400, detail=f"{i}번째 항목: 결제처가 비어 있습니다")
        if row["final_category"] not in user_categories:
            raise HTTPException(status_code=400,
                                detail=f"{i}번째 항목({row['merchant_name']}): 카테고리를 선택해 주세요")
        prepared.append(row)

    existing = _existing_keys(session, user_id, [r["import_key"] for r in prepared])
    suggestions = classify_merchants(session, user_id, {r["merchant_name"] for r in prepared})
    label = LABEL_TOSS if source_format == FORMAT_TOSS else LABEL_GENERIC
    touched_dates = set()
    inserted = skipped = 0

    for row in prepared:
        if row["import_key"] in existing:
            skipped += 1
            continue
        occurred_at = row["occurred_at"]
        suggested = suggestions.get(row["merchant_name"], (None, "NONE"))[0]
        corrected = suggested != row["final_category"]
        session.add(Transaction(
            user_id=user_id,
            tx_date=occurred_at.date(),
            tx_time=occurred_at.time().replace(microsecond=0),
            amount=int(row["amount"]),
            merchant_name=row["merchant_name"].strip()[:100],
            mydata_category=f"{label}:{row['tx_type']}" if row.get("tx_type") else label,
            final_category=row["final_category"],
            is_user_corrected=corrected,
            memo=(row.get("memo") or None),
            import_key=row["import_key"],
        ))
        if corrected:
            remember_user_choice(session, user_id, row["merchant_name"], row["final_category"])
        touched_dates.add(occurred_at.date())
        inserted += 1

    cancelled = 0
    for tx_id in dict.fromkeys(cancel_tx_ids or []):
        tx = session.get(Transaction, tx_id)
        # 이 유저가 예전에 파일로 가져온 거래만 취소할 수 있다
        if tx is None or tx.user_id != user_id or tx.import_key is None:
            continue
        touched_dates.add(tx.tx_date)
        session.delete(tx)
        cancelled += 1

    if touched_dates:
        session.execute(
            delete(GeneratedReport)
            .where(GeneratedReport.user_id == user_id)
            .where(GeneratedReport.period_start <= max(touched_dates))
            .where(GeneratedReport.period_end >= min(touched_dates))
        )

    try:
        session.commit()
    except IntegrityError:
        session.rollback()
        raise HTTPException(status_code=409, detail="같은 파일을 동시에 등록하고 있습니다. 잠시 후 다시 시도해 주세요")

    return {"inserted": inserted, "skipped_duplicate": skipped, "cancelled": cancelled,
            "dates": sorted(d.isoformat() for d in touched_dates)}
